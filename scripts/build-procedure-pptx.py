"""Build a Japanese operator procedure PPTX from captured screens."""
from __future__ import annotations

from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
SCREENS = ROOT / "docs" / "procedure" / "screens"
OUT = ROOT / "docs" / "procedure" / "UGV_MLOps_データ取込からデプロイ手順.pptx"

NAVY = RGBColor(0x0E, 0x14, 0x18)
PANEL = RGBColor(0x17, 0x20, 0x28)
TEAL = RGBColor(0x3D, 0xD6, 0xC6)
TEXT = RGBColor(0xE7, 0xEE, 0xF4)
MUTED = RGBColor(0x93, 0xA3, 0xB3)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK = RGBColor(0x1A, 0x24, 0x2C)
ACCENT2 = RGBColor(0x7E, 0xB6, 0xFF)
WARN = RGBColor(0xE0, 0xB1, 0x5B)

W = Inches(13.333)
H = Inches(7.5)
FONT = "Yu Gothic UI"


def set_font(run, size=16, bold=False, color=TEXT, name=FONT):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    latin = rPr.get_or_add_latin()
    latin.set("typeface", name)
    ns = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
    for tag in ("ea", "cs"):
        el = rPr.find(f"{ns}{tag}")
        if el is None:
            el = etree.SubElement(rPr, f"{ns}{tag}")
        el.set("typeface", name)


def add_run(tf, text, size=16, bold=False, color=TEXT):
    p = tf.paragraphs[0] if not tf.paragraphs[0].text else tf.add_paragraph()
    if not tf.paragraphs[0].text and tf.paragraphs[0] is p:
        run = p.add_run()
    else:
        if tf.paragraphs[0].text == "" and len(tf.paragraphs) == 1:
            run = p.add_run()
        else:
            run = p.add_run()
    run.text = text
    set_font(run, size=size, bold=bold, color=color)
    p.space_after = Pt(6)
    return p


def write_lines(shape, lines, size=16, color=TEXT):
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    first = True
    for item in lines:
        if isinstance(item, tuple):
            text, sz, bold, col = item
        else:
            text, sz, bold, col = item, size, False, color
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        run = p.add_run()
        run.text = text
        set_font(run, size=sz, bold=bold, color=col)
        p.space_after = Pt(8)
        p.level = 0


def fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def bar(slide, y=0):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, y, W, Inches(0.08))
    fill(s, TEAL)
    return s


def bg(slide, color=NAVY):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H)
    fill(s, color)
    spTree = slide.shapes._spTree
    sp = s._element
    spTree.remove(sp)
    spTree.insert(2, sp)


def title_box(slide, text, left=0.5, top=0.22, width=12.3, size=28):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(0.6))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = text
    set_font(run, size=size, bold=True, color=WHITE)
    return box


def caption(slide, text, left, top, width, height=0.4, size=12, color=MUTED):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = text
    set_font(run, size=size, color=color)
    return box


def picture(slide, name, left, top, width, height=None):
    path = SCREENS / name
    if not path.exists():
        box = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(left),
            Inches(top),
            Inches(width),
            Inches(height or 4.2),
        )
        fill(box, PANEL)
        write_lines(box, [f"画面未取得: {name}"], size=14, color=WARN)
        return box
    from PIL import Image

    with Image.open(path) as im:
        aspect = im.height / im.width if im.width else 0.6
    w = width
    h = w * aspect
    if height and h > height:
        h = height
        w = h / aspect if aspect else width
    return slide.shapes.add_picture(str(path), Inches(left), Inches(top), width=Inches(w), height=Inches(h))


def cover(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg(slide)
    bar(slide)
    title_box(slide, "UGV MLOps 操作手順書", top=2.1, size=40)
    box = slide.shapes.add_textbox(Inches(0.55), Inches(2.9), Inches(12), Inches(1.6))
    write_lines(
        box,
        [
            ("データ取り込みから車載デプロイまで", 22, False, TEAL),
            ("実際のポータル画面に沿った 8 工程", 18, False, MUTED),
            ("入口  http://127.0.0.1:8080", 16, False, ACCENT2),
        ],
    )
    cap = slide.shapes.add_textbox(Inches(0.55), Inches(6.7), Inches(12), Inches(0.4))
    write_lines(cap, [("オンプレ OSS ／ クラウド API なし ／ 機密サンプルは学習セットに入れない", 12, False, MUTED)])


def overview(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg(slide)
    bar(slide)
    title_box(slide, "全体の流れ（8 工程）")
    steps = [
        ("1", "帰還データを置く", "inbox フォルダへ bag / 画像"),
        ("2", "カタログへ取り込む", "パース・Garage 保管"),
        ("3", "ラベルを付ける", "LS / SALT / Xtreme1"),
        ("4", "データセットを切る", "知覚 / 検出 / E2E"),
        ("5", "学習する", "指標は MLflow"),
        ("6", "評価して通す", "ゲート + 署名"),
        ("7", "車載パッケージ", "S3 / S1 向け zip"),
        ("8", "UGV へ同期", "チャンク配送 → acked"),
    ]
    for i, (n, t, d) in enumerate(steps):
        col = i % 4
        row = i // 4
        x = 0.45 + col * 3.2
        y = 1.15 + row * 2.85
        card = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(3.0), Inches(2.5)
        )
        fill(card, PANEL)
        num = slide.shapes.add_textbox(Inches(x + 0.18), Inches(y + 0.2), Inches(2.6), Inches(0.4))
        write_lines(num, [(f"工程 {n}", 12, True, TEAL)])
        ttl = slide.shapes.add_textbox(Inches(x + 0.18), Inches(y + 0.65), Inches(2.6), Inches(0.9))
        write_lines(ttl, [(t, 18, True, WHITE)])
        dsc = slide.shapes.add_textbox(Inches(x + 0.18), Inches(y + 1.55), Inches(2.6), Inches(0.7))
        write_lines(dsc, [(d, 13, False, MUTED)])


def start_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg(slide)
    bar(slide)
    title_box(slide, "起動して入口を開く")
    left = slide.shapes.add_textbox(Inches(0.5), Inches(1.05), Inches(5.8), Inches(5.8))
    write_lines(
        left,
        [
            ("1. Docker Desktop を起動する", 16, True, WHITE),
            ("2. ランチャーで「中核スタック起動」または", 16, False, TEXT),
            ("    docker compose up -d --no-build --pull never", 14, False, TEAL),
            ("3. ブラウザは localhost ではなく 127.0.0.1 を使う", 16, False, TEXT),
            ("", 10, False, TEXT),
            ("ポータル　http://127.0.0.1:8080", 16, True, ACCENT2),
            ("Catalog　http://127.0.0.1:8000/docs", 14, False, MUTED),
            ("MLflow　　http://127.0.0.1:5000", 14, False, MUTED),
            ("Label Studio　http://127.0.0.1:8081", 14, False, MUTED),
            ("", 10, False, TEXT),
            ("初期ログイン（機密区画では必ず変更）", 16, True, WARN),
            ("Label Studio  admin@ugv.local / label_change_me", 14, False, TEXT),
        ],
    )
    picture(slide, "01_dashboard.png", 6.4, 1.1, 6.5, 5.6)


def shot_slide(prs, title, image, bullets, note=""):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg(slide)
    bar(slide)
    title_box(slide, title, width=12.4, size=24)
    left = slide.shapes.add_textbox(Inches(0.45), Inches(0.95), Inches(4.35), Inches(5.7))
    write_lines(left, [(b, 15, False, TEXT) for b in bullets])
    if note:
        cap = slide.shapes.add_textbox(Inches(0.45), Inches(6.55), Inches(4.35), Inches(0.7))
        write_lines(cap, [(note, 12, False, WARN)])
    picture(slide, image, 4.95, 0.95, 8.0, 5.95)


def wide_shot(prs, title, image, subtitle=""):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg(slide)
    bar(slide)
    title_box(slide, title, size=24)
    if subtitle:
        caption(slide, subtitle, 0.5, 0.78, 12.3, size=13, color=MUTED)
        picture(slide, image, 0.45, 1.2, 12.4, 5.85)
    else:
        picture(slide, image, 0.45, 1.0, 12.4, 6.1)


def ids_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg(slide)
    bar(slide)
    title_box(slide, "画面の場所と初期 ID")
    rows = [
        ("ポータル", "http://127.0.0.1:8080", "なし"),
        ("Catalog API", "http://127.0.0.1:8000/docs", "なし"),
        ("MLflow", "http://127.0.0.1:5000", "なし"),
        ("Label Studio", "http://127.0.0.1:8081", "admin@ugv.local / label_change_me"),
        ("Garage S3", "http://127.0.0.1:9000", "Access Key は .env（コンソールなし）"),
        ("Xtreme1", "http://127.0.0.1:8190", "初回 Sign up（本 compose 外）"),
        ("SALT", "WSL Ubuntu-24.04", "ログインなし"),
    ]
    y = 1.1
    hdr = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(y), Inches(12.3), Inches(0.38))
    fill(hdr, PANEL)
    for i, h in enumerate(("ツール", "URL / 場所", "ID / パスワード")):
        w = [2.4, 5.2, 4.7][i]
        x = 0.55 + sum([2.4, 5.2][:i])
        box = slide.shapes.add_textbox(Inches(x), Inches(y + 0.02), Inches(w), Inches(0.34))
        write_lines(box, [(h, 12, True, TEAL)])
    y = 1.52
    for tool, url, cred in rows:
        row = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(y), Inches(12.3), Inches(0.68))
        fill(row, PANEL if int(y * 10) % 2 == 0 else DARK)
        vals = (tool, url, cred)
        widths = (2.4, 5.2, 4.7)
        x = 0.55
        for v, w in zip(vals, widths):
            box = slide.shapes.add_textbox(Inches(x), Inches(y + 0.12), Inches(w), Inches(0.48))
            write_lines(box, [(v, 13, False, WHITE)])
            x += w
        y += 0.72
    caption(slide, "値は .env のデモ初期値です。機密区画・本番では起動前に必ず変更してください。", 0.5, 6.7, 12.3, color=WARN)


def notes_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg(slide)
    bar(slide)
    title_box(slide, "運用上の注意")
    items = [
        "知覚（分類・検出）はラベルなしでデータセットを切らない。E2E は cmd_vel が教師なので工程3をスキップ可。",
        "UNCLASSIFIED 以外のサンプルは解除するまで学習セットに入らない。",
        "公開 rosbag2 デモは短いクリップ確認用。ゲートは smoke。実任務モデルは staging / production。",
        "SALT / Xtreme1 は本 compose に載せない。工程3で作業ファイルを data/annotate に出してから各 GUI を使う。",
        "Xtreme1 の MinIO（:8193/:8194）と本スタックの Garage（:9000）は別。キーを混ぜない。",
        "ブラウザは http://127.0.0.1:... を使う（localhost だと IPv6 で失敗することがある）。",
        "失敗ジョブはダッシュボードの「最近の動き」か GET /jobs/{id} の error / reasons を見る。",
    ]
    box = slide.shapes.add_textbox(Inches(0.55), Inches(1.1), Inches(12.2), Inches(5.8))
    write_lines(box, [(f"・ {t}", 16, False, TEXT) for t in items])


def end_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg(slide)
    bar(slide)
    title_box(slide, "関連資料", top=2.4, size=32)
    box = slide.shapes.add_textbox(Inches(0.7), Inches(3.2), Inches(12), Inches(2.5))
    write_lines(
        box,
        [
            ("docs/usage.md　日常操作", 18, False, TEXT),
            ("docs/annotation.md　Label Studio / SALT / Xtreme1", 18, False, TEXT),
            ("docs/workflow.md　四脚 PJ との層対応", 18, False, TEXT),
            ("docs/setup-windows.md　新しい Windows PC への構築", 18, False, TEXT),
        ],
    )


def build() -> Path:
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    cover(prs)
    overview(prs)
    start_slide(prs)
    wide_shot(prs, "ダッシュボード — いまやること", "01_dashboard.png", "「いまやること」と工程カードから次の作業に進みます。")
    wide_shot(prs, "データタブ — 撮りためた中身を見る", "03_data_preview.png", "ミッションを選ぶとカメラ・cmd_vel / IMU・LiDAR（上から）が出ます。")
    wide_shot(prs, "8 工程フロー", "04_flow_overview.png", "上のカードを押すと、その工程の実行フォームが下に出ます。Swagger を開く必要はありません。")
    shot_slide(
        prs,
        "工程1  帰還データを置く",
        "05_step1_collect.png",
        [
            "ホストの data\\inbox\\<mission_id>\\ にコピーします。",
            "ROS2 bag（metadata.yaml + .db3 / .mcap）でも、展開済み画像でも可。",
            "ポータルでフォルダを再読込して、一覧に出たことを確認。",
            "サンプル bag は M-AVBAG-001（公開 CostNav を短く切ったもの）。",
            "確認できたら「取込へ進む」。",
        ],
        "inbox はまだ Catalog に入る前の置き場です。",
    )
    shot_slide(
        prs,
        "工程2  カタログへ取り込む",
        "06_step2_ingest.png",
        [
            "inbox のフォルダを選ぶ（または名前を入力）。",
            "種別は dropzone（自動判定）でよい。bag なら rosbag。",
            "「取込を実行」→ ジョブが completed になるまで待つ。",
            "中ではコピー → パース（JPEG / LiDAR / cmd_vel）→ DVC 追跡 → Garage 保管。",
            "データタブでカメラと点群が出れば成功。",
        ],
    )
    shot_slide(
        prs,
        "工程3  ラベルを付ける（ツールを選ぶ）",
        "07_step3_annotate.png",
        [
            "ミッションを選ぶ。",
            "アノテーションツールを選ぶ。",
            "  ・Label Studio … 画像クラス / bbox",
            "  ・SALT … LiDAR 点群（WSL GUI）",
            "  ・Xtreme1 … カメラと点群の確認",
            "「作業データを出す」で data\\annotate に書き出す。",
            "SALT / Xtreme1 の結果 jsonl は「結果を Catalog に戻す」。",
            "E2E（cmd_vel 教師）はスキップして工程4へ。",
        ],
        "知覚はラベル必須。未ラベルのまま obstacle_cls / detection は切れません。",
    )
    shot_slide(
        prs,
        "工程3  Label Studio（画像）",
        "14_step3_labelstudio.png",
        [
            "ツールに Label Studio を選んで投入する。",
            "http://127.0.0.1:8081 でログイン。",
            "プロジェクト UGV-obstacle-cls。",
            "クラスは clear / obstacle / person。",
            "データセット構築時に Catalog が export して annotation へ戻す。",
        ],
    )
    if (SCREENS / "16_labelstudio.png").exists():
        wide_shot(prs, "Label Studio 画面", "16_labelstudio.png", "admin@ugv.local / label_change_me（デモ初期値）")
    shot_slide(
        prs,
        "工程3  SALT（点群）",
        "07_step3_annotate.png",
        [
            "ツールに SALT を選ぶ。",
            "点群が data\\annotate\\<mission>\\lidar に .pcd で出る。",
            "WSL Ubuntu-24.04、conda 環境 SALT で GUI を開く。",
            "塗った結果を export.jsonl（各行 sample_id + label_json）に書く。",
            "工程3の「結果を Catalog に戻す」。",
        ],
    )
    shot_slide(
        prs,
        "工程3  Xtreme1（確認）",
        "13_step3_xtreme1.png",
        [
            "ツールに Xtreme1 を選ぶと画像と .pcd を出す。",
            "http://127.0.0.1:8190 を開く（Autoware compose。本スタック外）。",
            "初回は画面の Sign up。",
            "箱検出の確認用。Freespace マスクには使わない。",
            "書き出しがあれば export.jsonl を戻す。",
        ],
    )
    shot_slide(
        prs,
        "工程4  学習用データセットを切る",
        "08_step4_dataset.png",
        [
            "データセット ID と表示名を入れる。",
            "タスクを選ぶ。",
            "  ・obstacle_cls … 知覚分類（ラベル必須）",
            "  ・detection … YOLOX（bbox）",
            "  ・e2e … カメラ + cmd_vel（人手ラベル不要）",
            "「データセットを生成」。版 ID（例 ds-av-rosbag-v2）を控える。",
        ],
    )
    shot_slide(
        prs,
        "工程5  学習する",
        "09_step5_train.png",
        [
            "データセット版を選ぶ。",
            "学習段は前線 S3（既定）または後方 S4。",
            "「学習を実行」。run_id を控える。",
            "指標は MLflow http://127.0.0.1:5000",
            "  分類=精度、検出=mAP@0.5、E2E=mae_steer / mae_speed。",
        ],
    )
    if (SCREENS / "17_mlflow.png").exists():
        wide_shot(prs, "MLflow — 学習指標", "17_mlflow.png", "実験 ugv-obstacle-cls / ugv-e2e / 検出。run を開いて metrics を確認します。")
    shot_slide(
        prs,
        "工程6  評価して通す（ゲート）",
        "10_step6_gate.png",
        [
            "学習 run を選ぶ。",
            "ゲートは staging または production。",
            "「ゲートを実行」。通ると model_id と署名が付く。",
            "閾値は configs/artifact_policy.yaml。",
            "公開サンプルの短いクリップは smoke（デモ専用）。",
            "不合格なら reasons を読んでデータを足すか閾値を見直す。",
        ],
    )
    shot_slide(
        prs,
        "工程7  車載パッケージを作る",
        "11_step7_package.png",
        [
            "ゲート合格モデル（staging / canary / production）を選ぶ。",
            "対象段階 S3（前線）または S1（車載・小さい成果物のみ）。",
            "「パッケージを作る」。Garage packages に zip。",
            "フル重みは S3/S4。S1 はアダプタ等の小さいものだけ。",
        ],
    )
    shot_slide(
        prs,
        "工程8  UGV へ差分同期",
        "12_step8_sync.png",
        [
            "パッケージと車両 ID（例 UGV-01）を入れる。",
            "「同期を実行」。チャンク送信までこの画面で完了する。",
            "全部送れると acked。失敗時はロールバック。",
            "OS イメージ更新とは分け、モデル zip だけを送る。",
        ],
    )
    shot_slide(
        prs,
        "デプロイ後 — 課題を追跡する",
        "15_issues.png",
        [
            "課題タブで、デプロイ済みモデルに紐づけて起票。",
            "着手 → 解決（または対応しない / 再オープン）。",
            "ダッシュボードの「未解決課題」から同じ一覧に入れる。",
            "再現データは inbox に戻して工程2から再学習できる。",
        ],
    )
    wide_shot(
        prs,
        "一気通貫の確認（デモ）",
        "01_dashboard.png",
        "ポータル右上: 知覚デモ / E2Eデモ / YOLOXデモ / rosbag2デモ。rosbag2 は公開 CostNav を inbox の M-AVBAG-001 に載せて E2E まで回します。",
    )
    ids_slide(prs)
    notes_slide(prs)
    end_slide(prs)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUT))
    print("wrote", OUT, OUT.stat().st_size)
    return OUT


if __name__ == "__main__":
    build()
